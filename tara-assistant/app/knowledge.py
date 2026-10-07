"""Document parsing, cleaning, section-aware chunking and ingestion.

Formats: PDF, DOCX, XLSX / CSV (communication matrix, asset lists), JSON (threat library, prior TARA,
control catalogue), Markdown / text (architecture descriptions).
"""
import hashlib
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from . import config
from .taxonomy import detect_interfaces, normalise_names

IFACE_BY_NAME = {"can": "CAN", "ethernet": "ETH", "diagnostics": "DIAG", "diagnostic": "DIAG", "ota": "OTA", "backend": "BACKEND"}
HEADING_RE = re.compile(r"^(#{1,4}\s+.+|\d+(\.\d+)*\.?\s+[A-Z][^\n]{2,80})$")


@dataclass
class Chunk:
    id: str
    text: str
    meta: dict = field(default_factory=dict)


@dataclass
class ParsedDoc:
    document: str
    doc_type: str  # architecture | comm_matrix | asset_catalogue | threat_library | prior_tara | controls
    text: str = ""
    records: list = field(default_factory=list)
    version: str = ""
    review_date: str = ""
    raw_len: int = 0


def est_tokens(text: str) -> int:
    return int(len(text.split()) * 1.3)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


# ---------------------------------------------------------------- cleaning
def clean_text(text: str, repeated_lines: set | None = None) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    if repeated_lines:  # headers / footers detected across pages
        lines = [ln for ln in lines if ln.strip() not in repeated_lines]
    out, seen_para = [], set()
    for para in re.split(r"\n{2,}", "\n".join(lines)):
        p = re.sub(r"[ \t]+", " ", para).strip()
        if not p:
            continue
        key = re.sub(r"\W+", " ", p.lower())
        if key in seen_para and len(p) > 40:  # exact duplicate paragraph
            continue
        seen_para.add(key)
        out.append(p)
    return normalise_names("\n\n".join(out))


def _repeated_lines(pages: list[str]) -> set:
    if len(pages) < 3:
        return set()
    c = Counter()
    for p in pages:
        for ln in {l.strip() for l in p.split("\n") if l.strip()}:
            c[ln] += 1
    return {ln for ln, n in c.items() if n >= max(3, int(0.6 * len(pages))) and len(ln) < 120}


# ---------------------------------------------------------------- parsing
def parse_document(name: str, data: bytes, doc_type: str | None = None) -> ParsedDoc:
    ext = Path(name).suffix.lower()
    if ext in (".md", ".txt"):
        raw = data.decode("utf-8", "replace")
        d = ParsedDoc(Path(name).stem, doc_type or "architecture", text=clean_text(raw), raw_len=len(raw))
    elif ext == ".pdf":
        from pypdf import PdfReader

        pages = [(p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages]
        rep = _repeated_lines(pages)
        raw = "\n\n".join(pages)
        d = ParsedDoc(Path(name).stem, doc_type or "architecture", text=clean_text(raw, rep), raw_len=len(raw))
    elif ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = []
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            style = (p.style.name or "").lower()
            parts.append(("## " + p.text) if style.startswith("heading") else p.text)
        for t in doc.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        raw = "\n\n".join(parts)
        d = ParsedDoc(Path(name).stem, doc_type or "architecture", text=clean_text(raw), raw_len=len(raw))
    elif ext in (".csv", ".xlsx", ".xls"):
        df = pd.read_csv(io.BytesIO(data), dtype=str) if ext == ".csv" else pd.read_excel(io.BytesIO(data), dtype=str)
        df = df.fillna("").drop_duplicates()
        df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
        cols = set(df.columns)
        dt = doc_type or ("asset_catalogue" if "asset_id" in cols else "comm_matrix")
        recs = [{k: normalise_names(str(v).strip()) for k, v in r.items()} for r in df.to_dict("records")]
        rd = next((r.get("review_date") for r in recs if r.get("review_date")), "")
        d = ParsedDoc(Path(name).stem, dt, records=recs, review_date=rd, raw_len=len(data))
    elif ext == ".json":
        j = json.loads(data.decode("utf-8"))
        d = ParsedDoc(j.get("document", Path(name).stem), j["doc_type"], records=j["entries"], version=str(j.get("version", "")),
                      review_date=j.get("review_date", ""), raw_len=len(data))
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return d


# ---------------------------------------------------------------- chunking
def _split_sections(text: str) -> list[tuple[str, str]]:
    sections, title, buf = [], "Introduction", []
    for para in text.split("\n\n"):
        first = para.split("\n")[0].strip()
        if HEADING_RE.match(first) and len(para) < 5000:
            if buf:
                sections.append((title, "\n\n".join(buf)))
            title = re.sub(r"^#+\s*", "", first)
            buf = [para.split("\n", 1)[1]] if "\n" in para else []
        else:
            buf.append(para)
    if buf:
        sections.append((title, "\n\n".join(buf)))
    return [(t, b.strip()) for t, b in sections if b.strip()]


def _windows(text: str, max_tokens: int, overlap: float) -> list[str]:
    words = text.split()
    max_w = int(max_tokens / 1.3)
    if len(words) <= max_w:
        return [text]
    step = max(1, int(max_w * (1 - overlap)))
    out = []
    for i in range(0, len(words), step):
        out.append(" ".join(words[i:i + max_w]))
        if i + max_w >= len(words):
            break
    return out


def _section_interface(title: str, body: str) -> str:
    m = re.search(r"interface:?\s*([A-Za-z\-]+)", title, re.I)
    if m and m.group(1).lower() in IFACE_BY_NAME:
        return IFACE_BY_NAME[m.group(1).lower()]
    found = detect_interfaces(title)
    if found and not re.search(r"overview|trust", title, re.I):
        return found[0]
    return "GENERAL"


def make_citation(doc: ParsedDoc, record_id: str = "", location: str = "") -> str:
    left = f"{record_id} | " if record_id else ""
    return f"{left}{doc.document}{', ' + location if location else ''}"


def chunk_document(doc: ParsedDoc, project: str, sha: str) -> list[Chunk]:
    base = {"document": doc.document, "doc_type": doc.doc_type, "version": doc.version, "review_date": doc.review_date,
            "project": project, "sha": sha}
    chunks: list[Chunk] = []
    slug = re.sub(r"\W+", "-", doc.document)[:40]

    if doc.doc_type == "architecture":
        for si, (title, body) in enumerate(_split_sections(doc.text)):
            iface = _section_interface(title, body)
            for wi, w in enumerate(_windows(f"{title}\n{body}", config.CHUNK_MAX_TOKENS, config.CHUNK_OVERLAP)):
                meta = dict(base, interface=iface, asset_tag="", section=title, record_id=f"S{si + 1}", location=f"section '{title}'")
                meta["citation"] = make_citation(doc, location=f"section '{title}'")
                chunks.append(Chunk(f"{project}:{slug}:s{si}:{wi}", f"[{doc.document} | {title}]\n{w}", meta))
    elif doc.doc_type == "comm_matrix":
        for i, r in enumerate(doc.records):
            iface = r.get("interface") or (detect_interfaces(r.get("protocol", "") + " " + r.get("bus", "")) or ["GENERAL"])[0]
            rid = r.get("message") or f"row{i + 1}"
            text = f"[{doc.document} | communication matrix row {i + 1}] " + "; ".join(f"{k}: {v}" for k, v in r.items() if v)
            meta = dict(base, interface=iface, asset_tag="", section="communication matrix", record_id=rid, location=f"row {i + 1}")
            meta["citation"] = make_citation(doc, rid, f"row {i + 1}")
            chunks.append(Chunk(f"{project}:{slug}:r{i}", text, meta))
    elif doc.doc_type == "asset_catalogue":
        for i, r in enumerate(doc.records):
            rid = r.get("asset_id") or f"row{i + 1}"
            text = (f"[{doc.document} | asset {rid}] Asset {r.get('name', '')} ({r.get('type', '')}), interface {r.get('interface', '')}, "
                    f"properties {r.get('properties', '')}. {r.get('description', '')}")
            meta = dict(base, interface=r.get("interface", "GENERAL"), asset_tag=r.get("tag", ""), section="asset catalogue",
                        record_id=rid, location=f"row {i + 1}", asset_name=r.get("name", ""), asset_type=r.get("type", ""),
                        properties=r.get("properties", ""), owner=r.get("owner", ""), description=r.get("description", ""),
                        version=r.get("version", ""), review_date=r.get("review_date", doc.review_date))
            meta["citation"] = make_citation(doc, rid, f"row {i + 1}")
            chunks.append(Chunk(f"{project}:{slug}:{rid}", text, meta))
    elif doc.doc_type in ("threat_library", "prior_tara"):
        for e in doc.records:  # one chunk per entry so each scenario is retrieved and cited whole
            text = (f"[{doc.document} | {e['id']}] {e['title']}. {e['description']} Attack vector: {e['attack_vector']} "
                    f"Damage scenario: {e['damage_scenario']} Interface: {e['interface']}. STRIDE: {e['stride']}.")
            meta = dict(base, document=e.get("source_doc", doc.document), version=str(e.get("version", doc.version)),
                        review_date=e.get("review_date", doc.review_date), interface=e["interface"], asset_tag=e["asset_tag"],
                        section="threat entry", record_id=e["id"], location=e.get("location", ""), title=e["title"],
                        entry_json=json.dumps(e))
            meta["citation"] = f"{e['id']} | {meta['document']}" + (f", {e['location']}" if e.get("location") else "")
            chunks.append(Chunk(f"{project}:{e['id']}", text, meta))
    elif doc.doc_type == "controls":
        for e in doc.records:
            text = (f"[{doc.document} | {e['id']}] {e['title']}. {e['requirement']} Goal: {e['cybersecurity_goal']} "
                    f"Interface: {e['interface']}. Tests: {', '.join(e['test_ids'])}.")
            meta = dict(base, interface=e["interface"], asset_tag="", section="control", record_id=e["id"], location="",
                        title=e["title"], entry_json=json.dumps(e))
            meta["citation"] = f"{e['id']} | {doc.document} v{doc.version}"
            chunks.append(Chunk(f"{project}:{e['id']}", text, meta))
    else:
        raise ValueError(f"Unknown doc_type {doc.doc_type}")
    for c in chunks:  # chroma metadata must be primitives
        c.meta = {k: ("" if v is None else v) for k, v in c.meta.items()}
    return chunks


GLOBAL_TYPES = {"threat_library", "prior_tara", "controls"}


def ingest(store, name: str, data: bytes, project: str, doc_type: str | None = None) -> dict:
    """Parse, clean, chunk, embed and persist one document. Returns an ingestion-log record."""
    doc = parse_document(name, data, doc_type)
    scope = config.GLOBAL_SCOPE if doc.doc_type in GLOBAL_TYPES else project
    sha = _sha(data)
    store.delete(project=scope, document=doc.document)  # rebuild when approved knowledge changes
    chunks = chunk_document(doc, scope, sha)
    store.upsert(chunks)
    toks = [est_tokens(c.text) for c in chunks]
    ifaces = sorted({c.meta["interface"] for c in chunks if c.meta["interface"] not in ("", "GENERAL")})
    return {"document": doc.document, "doc_type": doc.doc_type, "scope": scope, "version": doc.version, "review_date": doc.review_date,
            "sha256": sha, "n_chunks": len(chunks), "avg_tokens": round(sum(toks) / len(toks), 1) if toks else 0,
            "max_tokens": max(toks) if toks else 0, "interfaces": ifaces, "bytes": len(data), "chars_removed": max(0, doc.raw_len - sum(len(c.text) for c in chunks))
            if doc.doc_type == "architecture" else 0}
